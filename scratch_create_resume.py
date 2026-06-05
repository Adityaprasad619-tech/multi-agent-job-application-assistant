import docx

doc = docx.Document()
doc.add_heading("Jane Doe", level=0)
doc.add_paragraph("Email: jane.doe@example.com")
doc.add_paragraph("Phone: 555-0199")
doc.add_paragraph("LinkedIn: linkedin.com/in/janedoe")
doc.add_paragraph("GitHub: github.com/janedoe")
doc.add_paragraph("Location: San Francisco, CA")

doc.add_heading("Technical Skills", level=1)
doc.add_paragraph("Python, SQL, Machine Learning, Streamlit, Git")

doc.add_heading("Experience", level=1)
p1 = doc.add_paragraph()
p1.add_run("Software Engineer at NASA\n").bold = True
p1.add_run("Developed python-based tools for processing satellite data and automated workflows using Machine Learning.")

doc.add_heading("Projects", level=1)
p2 = doc.add_paragraph()
p2.add_run("Multi-Agent Searcher\n").bold = True
p2.add_run("Built a job search application leveraging multiple AI agents for resume tailoring and ATS analysis.")

doc.add_heading("Education", level=1)
doc.add_paragraph("BS in Computer Science, Stanford University")

doc.add_heading("Achievements", level=1)
doc.add_paragraph("Recipient of NASA Innovation Award")

doc.add_heading("Certifications", level=1)
doc.add_paragraph("AWS Certified Solutions Architect")

doc.save("test_resume.docx")
print("Saved test_resume.docx successfully!")
